"""
Módulo de Generación y Exportación de Reportes Oficiales en Excel (openpyxl y pandas).

Implementa la exportación de la 'Planilla Oficial de Estadísticas Mensuales y Planta Docente'
con sus cuadros normativos correspondientes:
- Hoja 1: 'Estadística Mensual' (Planta Funcional, Cátedras, Situación de Revista y Cómputo Horario).
- Hoja 2: 'Cuadro DL-01' (Exclusivo Licencias por Enfermedad - RN-04, Oficina 1452 - RN-07, Año - RN-08).
- Hoja 3: 'Cuadro DL-02' (Inasistencias, Permisos y Otras Licencias - RN-05, RN-06: 1 día por falta).

Garantiza la preservación de estilos corporativos institucionales, fórmulas dinámicas y volcado
directo desde DataFrames de Pandas sin corromper el archivo base ni los formatos de la administración.
"""

import os
from datetime import datetime, date, timedelta
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from database.connection import get_db_session
from database.schema import Designation, Teacher, Subject, Career, Substitution
from models.attendance_processor import AttendanceProcessor
from models.normative import NormativeClassifier, DEFAULT_OFFICE_CODE


class ReportService:
    """
    Servicio de compilación y exportación de reportes institucionales en formato Excel.
    """

    # Paleta de colores institucional
    NAVY_PRIMARY = "1B365D"
    NAVY_ACCENT = "2C5282"
    GRAY_LIGHT = "F7FAFC"
    WHITE = "FFFFFF"
    BORDER_COLOR = "D2D6DC"

    @classmethod
    def get_processed_dataframes(cls, month: int, year: int = None, career_id: int = None) -> tuple[pd.DataFrame, pd.DataFrame]:
        """
        Retorna la tupla de DataFrames (df_dl01, df_dl02) procesados con Pandas.
        """
        return AttendanceProcessor.process_monthly_attendance(month=month, year=year, career_id=career_id)

    @classmethod
    def _create_styles(cls):
        """Genera el catálogo de fuentes, rellenos y bordes corporativos."""
        title_font = Font(name="Calibri", size=15, bold=True, color=cls.WHITE)
        sub_font = Font(name="Calibri", size=10, italic=True, color=cls.WHITE)
        header_font = Font(name="Calibri", size=10, bold=True, color=cls.WHITE)
        data_font = Font(name="Calibri", size=10)
        bold_font = Font(name="Calibri", size=10, bold=True)

        navy_fill = PatternFill(start_color=cls.NAVY_PRIMARY, end_color=cls.NAVY_PRIMARY, fill_type="solid")
        accent_fill = PatternFill(start_color=cls.NAVY_ACCENT, end_color=cls.NAVY_ACCENT, fill_type="solid")
        zebra_fill = PatternFill(start_color=cls.GRAY_LIGHT, end_color=cls.GRAY_LIGHT, fill_type="solid")

        thin_side = Side(border_style="thin", color=cls.BORDER_COLOR)
        border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

        return {
            "title_font": title_font,
            "sub_font": sub_font,
            "header_font": header_font,
            "data_font": data_font,
            "bold_font": bold_font,
            "navy_fill": navy_fill,
            "accent_fill": accent_fill,
            "zebra_fill": zebra_fill,
            "border": border
        }

    @classmethod
    def dump_dataframe_to_sheet(
        cls,
        ws,
        df: pd.DataFrame,
        title_text: str,
        subtitle_text: str,
        total_label: str = "TOTAL DÍAS",
        days_col_idx: int = 11,
        total_cols: int = 12
    ):
        """
        Vuelca un DataFrame procesado de Pandas directamente en la hoja especificada de openpyxl,
        preservando o aplicando formatos oficiales, encabezados estilizados, bordes y fórmulas de totales.
        """
        styles = cls._create_styles()

        # 1. Título y Subtítulo
        last_col_letter = get_column_letter(total_cols)
        ws.merge_cells(f"A1:{last_col_letter}1")
        top_title = ws["A1"]
        top_title.value = title_text
        top_title.font = styles["title_font"]
        top_title.fill = styles["navy_fill"]
        top_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 34

        ws.merge_cells(f"A2:{last_col_letter}2")
        sub_title = ws["A2"]
        sub_title.value = subtitle_text
        sub_title.font = styles["sub_font"]
        sub_title.fill = styles["accent_fill"]
        sub_title.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[2].height = 22

        ws.row_dimensions[3].height = 10  # Espacio en blanco

        # 2. Encabezados de columnas (Fila 4)
        header_row = 4
        ws.row_dimensions[header_row].height = 26
        headers = list(df.columns)
        for col_idx, h in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h)
            cell.font = styles["header_font"]
            cell.fill = styles["navy_fill"]
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = styles["border"]

        # 3. Filas de datos (Fila 5 en adelante)
        start_row = 5
        rows_count = len(df)

        if rows_count == 0:
            # Caso sin novedades: Fila informativa sin romper fórmulas
            cur_row = start_row
            ws.row_dimensions[cur_row].height = 22
            ws.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=total_cols)
            msg_cell = ws.cell(row=cur_row, column=1, value="Sin registros de novedades para el período analizado.")
            msg_cell.font = Font(name="Calibri", size=10, italic=True, color="718096")
            msg_cell.alignment = Alignment(horizontal="center", vertical="center")
            for c_idx in range(1, total_cols + 1):
                ws.cell(row=cur_row, column=c_idx).border = styles["border"]

            tot_row = start_row + 1
            ws.row_dimensions[tot_row].height = 24
            ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=days_col_idx - 1)
            lbl_tot = ws.cell(row=tot_row, column=1, value=total_label)
            lbl_tot.font = styles["bold_font"]
            lbl_tot.alignment = Alignment(horizontal="right", vertical="center")

            c_val = ws.cell(row=tot_row, column=days_col_idx, value=0)
            c_val.font = styles["bold_font"]
            c_val.alignment = Alignment(horizontal="right", vertical="center")

            for c_idx in range(1, total_cols + 1):
                ws.cell(row=tot_row, column=c_idx).border = styles["border"]
                ws.cell(row=tot_row, column=c_idx).fill = styles["zebra_fill"]
        else:
            for i, row in enumerate(df.itertuples(index=False)):
                cur_row = start_row + i
                ws.row_dimensions[cur_row].height = 20

                for col_idx, val in enumerate(row, start=1):
                    cell = ws.cell(row=cur_row, column=col_idx, value=val)
                    cell.font = styles["data_font"]
                    cell.border = styles["border"]

                    # Formato explícito de texto para DNI (RN-01: preserva ceros y evita notación científica)
                    if headers[col_idx - 1] == "DNI":
                        cell.number_format = "@"
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif headers[col_idx - 1] in ("Oficina", "Año", "Mes", "Situación de Revista", "Desde", "Hasta"):
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif headers[col_idx - 1] in ("Total Días", "Hs Semanales", "Hs Mes"):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0"
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                    # Cebra en filas pares
                    if i % 2 == 1:
                        cell.fill = styles["zebra_fill"]

            # 4. Fila de Totales con fórmula oficial de administración openpyxl
            tot_row = start_row + rows_count
            ws.row_dimensions[tot_row].height = 24
            ws.merge_cells(start_row=tot_row, start_column=1, end_row=tot_row, end_column=days_col_idx - 1)
            lbl_tot = ws.cell(row=tot_row, column=1, value=total_label)
            lbl_tot.font = styles["bold_font"]
            lbl_tot.alignment = Alignment(horizontal="right", vertical="center")

            days_col_letter = get_column_letter(days_col_idx)
            # Fórmula oficial en openpyxl que se evalúa automáticamente en Excel
            c_tot = ws.cell(
                row=tot_row,
                column=days_col_idx,
                value=f"=SUM({days_col_letter}{start_row}:{days_col_letter}{tot_row - 1})"
            )
            c_tot.font = styles["bold_font"]
            c_tot.alignment = Alignment(horizontal="right", vertical="center")
            c_tot.number_format = "#,##0"

            for c_idx in range(1, total_cols + 1):
                ws.cell(row=tot_row, column=c_idx).border = styles["border"]
                ws.cell(row=tot_row, column=c_idx).fill = styles["zebra_fill"]

        # 5. Ajuste dinámico de ancho de columnas
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row in [1, 2]:
                    continue
                if cell.value:
                    val_str = str(cell.value)
                    # No medir fórmulas para el ancho
                    if not val_str.startswith("="):
                        max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    @classmethod
    def generate_monthly_statistics_excel(
        cls,
        month: int,
        year: int,
        career_id: int = None,
        output_path: str = None,
        base_template_path: str = None
    ) -> str:
        """
        Genera la 'Planilla Oficial de Estadísticas Mensuales y Asistencia Docente'
        respetando el formato institucional de educación superior.
        Incluye de forma automatizada:
        - Hoja 1: Estadística Mensual y Planta Funcional
        - Hoja 2: Cuadro DL-01 (Licencias por Enfermedad - RN-04, RN-07, RN-08)
        - Hoja 3: Cuadro DL-02 (Inasistencias y Permisos - RN-05, RN-06)
        """
        if not output_path:
            reports_dir = os.path.join(os.getcwd(), "reportes")
            os.makedirs(reports_dir, exist_ok=True)
            output_path = os.path.join(reports_dir, f"Planilla_Estadistica_Mensual_{year}_{month:02d}.xlsx")

        session = get_db_session()
        try:
            # 1. Consultar designaciones para Hoja 1
            query = session.query(Designation).join(Subject).join(Teacher).outerjoin(Career, Subject.career_id == Career.id)
            if career_id and career_id != 0:
                query = query.filter(Subject.career_id == career_id)

            designations = query.filter(Designation.is_active == True).order_by(
                Teacher.last_name, Designation.day_of_week
            ).all()

            # Estructurar datos de Hoja 1
            first_day_of_month = date(year, month, 1)
            if month == 12:
                last_day_of_month = date(year, 12, 31)
            else:
                last_day_of_month = date(year, month + 1, 1) - timedelta(days=1)

            rows_stats = []
            for d in designations:
                t = d.teacher
                s = d.subject
                c = s.career if s else None

                active_sub = session.query(Substitution).filter(
                    Substitution.designation_id == d.id,
                    Substitution.is_active == True,
                    Substitution.start_date <= last_day_of_month,
                    Substitution.end_date >= first_day_of_month
                ).first()

                condicion = "Titular en Ejercicio"
                detalle_suplencia = "-"
                resolucion = d.resolution_number or "-"

                if active_sub:
                    condicion = f"Licencia (Reemplazado por {active_sub.suplente.full_name})"
                    detalle_suplencia = f"{active_sub.license_reason} ({active_sub.start_date.strftime('%d/%m')} al {active_sub.end_date.strftime('%d/%m')})"
                    resolucion = active_sub.resolution_number

                horas_semanales = d.hours_count
                horas_mensuales = int(horas_semanales * 4)

                rows_stats.append({
                    "Legajo": t.legajo,
                    "Apellido y Nombres": t.full_name,
                    "DNI": NormativeClassifier.normalize_dni(t.dni),
                    "Situación de Revista": d.effective_revista,
                    "Carrera": c.name if c else "Sin Asignar",
                    "Cátedra / Espacio Curricular": d.effective_cargo,
                    "Día": d.day_of_week,
                    "Horario": f"{d.start_time} - {d.end_time}",
                    "Hs Cátedra Sem.": horas_semanales,
                    "Hs Cátedra Mes": horas_mensuales,
                    "Condición": condicion,
                    "N° Resolución Ministerial": resolucion,
                    "Motivo Licencia / Suplencia": detalle_suplencia
                })

            # 2. Procesar asistencias con Pandas para DL-01 y DL-02
            df_dl01, df_dl02 = AttendanceProcessor.process_monthly_attendance(
                month=month,
                year=year,
                career_id=career_id,
                session=session
            )

            # 3. Inicializar Workbook openpyxl
            styles = cls._create_styles()
            if base_template_path and os.path.exists(base_template_path):
                wb = load_workbook(base_template_path)
            else:
                wb = Workbook()

            # --- HOJA 1: Estadística Mensual y Planta Docente ---
            ws_stats = wb.active
            ws_stats.title = f"Estadística {month:02d}-{year}"

            ws_stats.merge_cells("A1:M1")
            top_title = ws_stats["A1"]
            top_title.value = "SIGET - SISTEMA INTEGRAL DE GESTIÓN EDUCATIVA TERCIARIA"
            top_title.font = styles["title_font"]
            top_title.fill = styles["navy_fill"]
            top_title.alignment = Alignment(horizontal="center", vertical="center")
            ws_stats.row_dimensions[1].height = 36

            ws_stats.merge_cells("A2:M2")
            sub_title = ws_stats["A2"]
            sub_title.value = f"PLANILLA OFICIAL DE ESTADÍSTICAS MENSUALES Y PLANTA FUNCIONAL - PERÍODO: {month:02d}/{year} (Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')})"
            sub_title.font = styles["sub_font"]
            sub_title.fill = styles["accent_fill"]
            sub_title.alignment = Alignment(horizontal="center", vertical="center")
            ws_stats.row_dimensions[2].height = 24

            ws_stats.row_dimensions[3].height = 10

            headers_stats = [
                "Legajo", "Apellido y Nombres", "DNI", "Situación Revista",
                "Carrera", "Cátedra / Espacio Curricular", "Día", "Horario",
                "Hs Semanales", "Hs Mes", "Condición en Planta",
                "N° Resolución", "Detalle Suplencia / Licencia"
            ]
            header_row = 4
            ws_stats.row_dimensions[header_row].height = 26
            for col_idx, h in enumerate(headers_stats, start=1):
                cell = ws_stats.cell(row=header_row, column=col_idx, value=h)
                cell.font = styles["header_font"]
                cell.fill = styles["navy_fill"]
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                cell.border = styles["border"]

            start_row = 5
            for i, r in enumerate(rows_stats):
                cur_row = start_row + i
                ws_stats.row_dimensions[cur_row].height = 20
                values = [
                    r["Legajo"], r["Apellido y Nombres"], r["DNI"], r["Situación de Revista"],
                    r["Carrera"], r["Cátedra / Espacio Curricular"], r["Día"], r["Horario"],
                    r["Hs Cátedra Sem."], r["Hs Cátedra Mes"], r["Condición"],
                    r["N° Resolución Ministerial"], r["Motivo Licencia / Suplencia"]
                ]
                for col_idx, val in enumerate(values, start=1):
                    cell = ws_stats.cell(row=cur_row, column=col_idx, value=val)
                    cell.font = styles["data_font"]
                    cell.border = styles["border"]

                    if col_idx == 3:  # DNI como texto
                        cell.number_format = "@"
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx in [1, 7, 8, 12]:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx in [9, 10]:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = "#,##0"
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                    if i % 2 == 1:
                        cell.fill = styles["zebra_fill"]

            # Fila de Totales Hoja 1
            tot_row_stats = start_row + len(rows_stats)
            ws_stats.row_dimensions[tot_row_stats].height = 24
            ws_stats.merge_cells(start_row=tot_row_stats, start_column=1, end_row=tot_row_stats, end_column=8)
            lbl_tot = ws_stats.cell(row=tot_row_stats, column=1, value="TOTALES GENERALES")
            lbl_tot.font = styles["bold_font"]
            lbl_tot.alignment = Alignment(horizontal="right", vertical="center")

            c_tot_sem = ws_stats.cell(
                row=tot_row_stats, column=9,
                value=f"=SUM(I{start_row}:I{tot_row_stats - 1})" if len(rows_stats) > 0 else 0
            )
            c_tot_sem.font = styles["bold_font"]
            c_tot_sem.alignment = Alignment(horizontal="right", vertical="center")
            c_tot_sem.number_format = "#,##0"

            c_tot_mes = ws_stats.cell(
                row=tot_row_stats, column=10,
                value=f"=SUM(J{start_row}:J{tot_row_stats - 1})" if len(rows_stats) > 0 else 0
            )
            c_tot_mes.font = styles["bold_font"]
            c_tot_mes.alignment = Alignment(horizontal="right", vertical="center")
            c_tot_mes.number_format = "#,##0"

            for col_idx in range(1, 14):
                ws_stats.cell(row=tot_row_stats, column=col_idx).border = styles["border"]
                ws_stats.cell(row=tot_row_stats, column=col_idx).fill = styles["zebra_fill"]

            for col in ws_stats.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row in [1, 2]:
                        continue
                    if cell.value:
                        v_str = str(cell.value)
                        if not v_str.startswith("="):
                            max_len = max(max_len, len(v_str))
                ws_stats.column_dimensions[col_letter].width = max(max_len + 4, 12)

            # --- HOJA 2: Cuadro DL-01 (Licencias por Enfermedad - RN-04) ---
            ws_dl01 = wb.create_sheet(title="Cuadro DL-01")
            cls.dump_dataframe_to_sheet(
                ws=ws_dl01,
                df=df_dl01,
                title_text="SIGET - CUADRO DL-01: LICENCIAS POR ENFERMEDAD (OFICINA 1452)",
                subtitle_text=f"Período: {month:02d}/{year} - Licencias Médicas, Art. 48, Salud Docente (Normativa RN-04 / RN-07: Oficina 1452 / RN-08: Año {year})",
                total_label="TOTAL DÍAS LICENCIA MÉDICA",
                days_col_idx=11,
                total_cols=12
            )

            # --- HOJA 3: Cuadro DL-02 (Inasistencias y Otras Licencias - RN-05 & RN-06) ---
            ws_dl02 = wb.create_sheet(title="Cuadro DL-02")
            cls.dump_dataframe_to_sheet(
                ws=ws_dl02,
                df=df_dl02,
                title_text="SIGET - CUADRO DL-02: INASISTENCIAS, PERMISOS Y OTRAS LICENCIAS (OFICINA 1452)",
                subtitle_text=f"Período: {month:02d}/{year} - Inasistencias de 1 día (RN-06), Permisos y Otras Licencias no Médicas (RN-05 / RN-07 / RN-08)",
                total_label="TOTAL DÍAS INASISTENCIAS Y PERMISOS",
                days_col_idx=11,
                total_cols=12
            )

            # Guardar archivo completo
            wb.save(output_path)
            return output_path

        finally:
            session.close()

    @classmethod
    def export_dl01_excel(cls, month: int, year: int, career_id: int = None, output_path: str = None) -> str:
        """Exporta exclusivamente el Cuadro DL-01 en un archivo independiente si es requerido."""
        if not output_path:
            reports_dir = os.path.join(os.getcwd(), "reportes")
            os.makedirs(reports_dir, exist_ok=True)
            output_path = os.path.join(reports_dir, f"Cuadro_DL01_Enfermedad_{year}_{month:02d}.xlsx")

        df_dl01, _ = AttendanceProcessor.process_monthly_attendance(month=month, year=year, career_id=career_id)
        wb = Workbook()
        ws = wb.active
        ws.title = "Cuadro DL-01"
        cls.dump_dataframe_to_sheet(
            ws=ws,
            df=df_dl01,
            title_text="SIGET - CUADRO DL-01: LICENCIAS POR ENFERMEDAD (OFICINA 1452)",
            subtitle_text=f"Período: {month:02d}/{year} - Licencias Médicas (RN-04 / RN-07 / RN-08)",
            total_label="TOTAL DÍAS LICENCIA MÉDICA",
            days_col_idx=11,
            total_cols=12
        )
        wb.save(output_path)
        return output_path

    @classmethod
    def export_dl02_excel(cls, month: int, year: int, career_id: int = None, output_path: str = None) -> str:
        """Exporta exclusivamente el Cuadro DL-02 en un archivo independiente si es requerido."""
        if not output_path:
            reports_dir = os.path.join(os.getcwd(), "reportes")
            os.makedirs(reports_dir, exist_ok=True)
            output_path = os.path.join(reports_dir, f"Cuadro_DL02_Inasistencias_{year}_{month:02d}.xlsx")

        _, df_dl02 = AttendanceProcessor.process_monthly_attendance(month=month, year=year, career_id=career_id)
        wb = Workbook()
        ws = wb.active
        ws.title = "Cuadro DL-02"
        cls.dump_dataframe_to_sheet(
            ws=ws,
            df=df_dl02,
            title_text="SIGET - CUADRO DL-02: INASISTENCIAS Y PERMISOS (OFICINA 1452)",
            subtitle_text=f"Período: {month:02d}/{year} - Inasistencias y Permisos (RN-05 / RN-06: 1 día por inasistencia)",
            total_label="TOTAL DÍAS INASISTENCIAS Y PERMISOS",
            days_col_idx=11,
            total_cols=12
        )
        wb.save(output_path)
        return output_path
