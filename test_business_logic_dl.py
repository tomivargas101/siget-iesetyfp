"""
Suite de Validación Exhaustiva de Lógica de Negocio y Reglas Normativas (RN-01 a RN-08).
Verifica:
- RN-01: Identificación unívoca por DNI como texto y almacenamiento único de datos comunes.
- RN-02: Múltiples designaciones por persona con situación de revista (TIT, INT, SUP).
- RN-03: Conservación del historial completo ante bajas lógicas (sin borrado físico).
- RN-04: Cuadro DL-01 exclusivo para licencias por enfermedad.
- RN-05: Cuadro DL-02 para inasistencias, permisos y licencias no médicas.
- RN-06: Una inasistencia solo puede abarcar un día (Desde == Hasta, Total Días == 1).
- RN-07: Autocompletado de Oficina: valor por defecto 1452.
- RN-08: Autocompletado de Año: año calendario actual.
- Agrupación automática de días consecutivos en licencias continuas.
- Exportación en openpyxl con fórmulas y preservación de formatos.
"""

import sys
from pathlib import Path
from datetime import date, timedelta
import pandas as pd
import openpyxl

BASE_DIR = Path(__file__).resolve().parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

from database.config import load_config
from database.connection import init_db, get_db_session
from database.schema import Teacher, Subject, Designation, Attendance, Substitution, Career
from models.normative import (
    NormativeClassifier,
    DEFAULT_OFFICE_CODE,
    REVISTA_TIT,
    REVISTA_INT,
    REVISTA_SUP
)
from models.teacher import TeacherService
from models.designation import DesignationService
from models.attendance import AttendanceService
from models.attendance_processor import AttendanceProcessor
from models.reports import ReportService


def run_all_normative_tests():
    print("======================================================================")
    print("INICIANDO SUITE DE TESTS: REGLAS NORMATIVAS INSTITUCIONALES (RN-01 A RN-08)")
    print("======================================================================\n")

    cfg = load_config()
    init_db(cfg)
    session = get_db_session()

    try:
        # ----------------------------------------------------------------------
        # TEST 1: REGLAS RN-01 Y RN-02 (Gestión de Personal y Designaciones Múltiples)
        # ----------------------------------------------------------------------
        print("--- [TEST 1] RN-01 & RN-02: DNI como Texto y Múltiples Designaciones ---")
        test_dni = "09456789"  # Con cero a la izquierda para validar tipado estricto como texto

        # 1.1 Guardar docente nuevo con DNI de texto
        ok, msg = TeacherService.save({
            "dni": test_dni,
            "legajo": "DOC-NORM-01",
            "cuil": "20-09456789-2",
            "first_name": "Mariana Andrea",
            "last_name": "Sarmiento",
            "email": "mariana.sarmiento@siget-edu.ar",
            "revista": "Titular",
            "max_weekly_hours": 30
        })
        assert ok or "Ya existe" in msg, f"Error al guardar docente: {msg}"

        docente = TeacherService.get_by_dni(test_dni)
        assert docente is not None, "El docente no fue encontrado por su DNI (RN-01)"
        assert docente.dni == test_dni, f"DNI corrupto: esperado {test_dni}, obtenido {docente.dni}"
        assert isinstance(docente.dni, str), "El DNI debe ser almacenado estrictamente como texto (RN-01)"
        print(f"  [OK] RN-01: Docente creado con DNI como texto: '{docente.dni}' - {docente.full_name}")

        # 1.2 Validar que NO permite duplicar a la persona por DNI (datos comunes se almacenan una sola vez)
        ok_dup, msg_dup = TeacherService.save({
            "dni": test_dni,
            "legajo": "DOC-OTRO-99",
            "first_name": "Mariana Copia",
            "last_name": "Sarmiento"
        })
        assert not ok_dup, "Debe rechazar la creación de otra persona con el mismo DNI (RN-01)"
        print(f"  [OK] RN-01: Rechazo de DNI duplicado verificado correctamente: '{msg_dup}'")

        # 1.3 Asignar múltiples designaciones con situaciones de revista diferenciadas (RN-02)
        mats = session.query(Subject).all()
        assert len(mats) >= 3, "Se requieren al menos 3 materias de prueba"
        m1, m2, m3 = mats[0], mats[1], mats[2]

        # Designación 1: Titular (TIT)
        desig1_data = {
            "teacher_id": docente.id,
            "subject_id": m1.id,
            "cargo": "Profesor de Cátedra",
            "revista": "TIT",
            "day_of_week": "Martes",
            "start_time": "08:00",
            "end_time": "10:00",
            "hours_count": 2,
            "resolution_number": "RES-TIT-101/26"
        }
        ok_d1, _ = DesignationService.save_designation(desig1_data)

        # Designación 2: Interino (INT)
        desig2_data = {
            "teacher_id": docente.id,
            "subject_id": m2.id,
            "cargo": "Jefe de Trabajos Prácticos",
            "revista": "INT",
            "day_of_week": "Miércoles",
            "start_time": "14:00",
            "end_time": "16:00",
            "hours_count": 2,
            "resolution_number": "RES-INT-202/26"
        }
        ok_d2, _ = DesignationService.save_designation(desig2_data)

        # Designación 3: Suplente (SUP)
        desig3_data = {
            "teacher_id": docente.id,
            "subject_id": m3.id,
            "cargo": "Profesor Reemplazante",
            "revista": "SUP",
            "day_of_week": "Jueves",
            "start_time": "18:00",
            "end_time": "20:00",
            "hours_count": 2,
            "resolution_number": "RES-SUP-303/26"
        }
        ok_d3, _ = DesignationService.save_designation(desig3_data)

        # Verificar designaciones vinculadas al DNI del docente (RN-01 y RN-02)
        docente_desigs = DesignationService.get_by_teacher_dni(test_dni)
        assert len(docente_desigs) >= 3, f"Faltan designaciones vinculadas al DNI: {len(docente_desigs)}"
        revistas = [d["revista"] for d in docente_desigs]
        assert "TIT" in revistas and "INT" in revistas and "SUP" in revistas, f"Faltan revistas TIT/INT/SUP: {revistas}"
        print(f"  [OK] RN-02: Múltiples designaciones confirmadas para '{test_dni}' con revistas {revistas}\n")

        # ----------------------------------------------------------------------
        # TEST 2: REGLA RN-03 (Historial Completo y Bajas Lógicas)
        # ----------------------------------------------------------------------
        print("--- [TEST 2] RN-03: Conservación del Historial Completo ante Bajas Lógicas ---")
        first_desig_id = docente_desigs[0]["id"]

        # Baja lógica de la designación
        ok_toggle, msg_toggle = DesignationService.toggle_active(first_desig_id)
        assert ok_toggle, f"Error al alternar estado: {msg_toggle}"
        print(f"  [OK] Designación {first_desig_id} dada de baja lógica: {msg_toggle}")

        # Comprobar que sigue existiendo en el historial
        history = DesignationService.get_history(teacher_id=docente.id)
        history_ids = [h["id"] for h in history]
        assert first_desig_id in history_ids, "La designación dada de baja lógica no debe ser eliminada del historial (RN-03)"

        # Comprobar historial general del docente
        full_hist = TeacherService.get_full_history(docente.id)
        assert full_hist["designations_count"] >= 3, "El contador histórico debe conservar todas las designaciones"
        print(f"  [OK] RN-03: Historial completo conservado íntegramente ({len(history)} designaciones registradas)\n")

        # ----------------------------------------------------------------------
        # TEST 3: REGLAS RN-04, RN-05, RN-06 Y AGRUPACIÓN DE DÍAS CONSECUTIVOS
        # ----------------------------------------------------------------------
        print("--- [TEST 3] RN-04, RN-05, RN-06 & Agrupación de Días Consecutivos ---")
        # Preparar registros de asistencia de prueba para el mes 10/2026:
        # Caso A: Licencia médica continua (5 días: 05/10/2026 a 09/10/2026) -> DL-01 (RN-04)
        med_start = date(2026, 10, 5)   # Lunes
        med_end = date(2026, 10, 9)     # Viernes
        ok_range, msg_range, count_range = AttendanceService.register_range(
            teacher_id=docente.id,
            start_date=med_start,
            end_date=med_end,
            status="Licencia",
            hours=4,
            observations="Art. 48 Tratamiento Médico Prolongado"
        )
        assert ok_range, f"Error al registrar rango médico: {msg_range}"
        assert count_range == 5, f"Se esperaban 5 días continuos, se registraron {count_range}"

        # Caso B: Inasistencias consecutivas (3 días: 13/10, 14/10, 15/10) -> DL-02 (RN-05 & RN-06)
        # Regla RN-06: Cada inasistencia solo puede abarcar 1 día (NUNCA deben agruparse)
        for offset in range(3):
            att_date = date(2026, 10, 13) + timedelta(days=offset)
            AttendanceService.save_record(
                teacher_id=docente.id,
                target_date=att_date,
                status="Ausente c/aviso",
                hours=4,
                observations="Falta justificada con aviso"
            )

        # Caso C: Otra licencia no médica (2 días continuos: 22/10 y 23/10) -> DL-02 (RN-05)
        ok_perm, _, _ = AttendanceService.register_range(
            teacher_id=docente.id,
            start_date=date(2026, 10, 22),
            end_date=date(2026, 10, 23),
            status="Licencia",
            hours=4,
            observations="Mesa de Examen / Estudio"
        )
        assert ok_perm

        # Procesar con Pandas
        df_dl01, df_dl02 = AttendanceProcessor.process_monthly_attendance(month=10, year=2026)

        print(f"  Resultados del Procesamiento con Pandas:")
        print(f"  - Filas Cuadro DL-01: {len(df_dl01)}")
        print(f"  - Filas Cuadro DL-02: {len(df_dl02)}")

        # 3.1 Validaciones en Cuadro DL-01 (Licencias por Enfermedad - RN-04)
        doc_dl01 = df_dl01[df_dl01["DNI"] == test_dni]
        assert not doc_dl01.empty, "El docente debe tener registro de licencia médica en DL-01 (RN-04)"
        rec_med = doc_dl01.iloc[0]
        assert rec_med["Desde"] == "05/10/2026", f"Fecha 'Desde' incorrecta: {rec_med['Desde']}"
        assert rec_med["Hasta"] == "09/10/2026", f"Fecha 'Hasta' incorrecta: {rec_med['Hasta']}"
        assert rec_med["Total Días"] == 5, f"Cómputo de días incorrecto: {rec_med['Total Días']}"
        assert "Art. 48" in str(rec_med["Artículo / Motivo"]) or "Salud" in str(rec_med["Artículo / Motivo"]), "Motivo incorrecto en DL-01"
        print(f"  [OK] RN-04: Cuadro DL-01 agrupó correctamente 5 días continuos: Desde {rec_med['Desde']} hasta {rec_med['Hasta']} (Total: {rec_med['Total Días']} días)")

        # 3.2 Validaciones en Cuadro DL-02 (RN-05 y RN-06)
        doc_dl02 = df_dl02[df_dl02["DNI"] == test_dni]
        assert not doc_dl02.empty, "El docente debe tener registros en DL-02 (RN-05)"

        # Comprobar RN-06: Inasistencias deben ser exactamente de 1 día cada una
        ina_records = doc_dl02[doc_dl02["Tipo Novedad / Motivo"].str.contains("Ausente|Falta", case=False, na=False)]
        assert len(ina_records) == 3, f"RN-06 violada: Se esperaban 3 filas individuales de inasistencia, se obtuvieron {len(ina_records)}"
        for _, r_ina in ina_records.iterrows():
            assert r_ina["Desde"] == r_ina["Hasta"], f"RN-06 violada: Desde ({r_ina['Desde']}) != Hasta ({r_ina['Hasta']})"
            assert r_ina["Total Días"] == 1, f"RN-06 violada: Total Días debe ser 1, obtenido {r_ina['Total Días']}"
        print(f"  [OK] RN-06: Verificado estrictamente: 3 inasistencias generaron 3 registros independientes de 1 día (Desde == Hasta)")

        # Comprobar licencia no médica continua agrupada en DL-02
        perm_records = doc_dl02[doc_dl02["Tipo Novedad / Motivo"].str.contains("Examen|Estudio", case=False, na=False)]
        assert not perm_records.empty, "Falta la licencia por Examen en DL-02"
        rec_perm = perm_records.iloc[0]
        assert rec_perm["Desde"] == "22/10/2026" and rec_perm["Hasta"] == "23/10/2026"
        assert rec_perm["Total Días"] == 2
        print(f"  [OK] RN-05: Licencia no médica agrupó 2 días continuos en DL-02 ({rec_perm['Desde']} al {rec_perm['Hasta']})\n")

        # ----------------------------------------------------------------------
        # TEST 4: REGLAS RN-07 Y RN-08 (Autocompletado Normativo)
        # ----------------------------------------------------------------------
        print("--- [TEST 4] RN-07 & RN-08: Autocompletado de Oficina 1452 y Año Actual ---")
        for df, name in [(df_dl01, "DL-01"), (df_dl02, "DL-02")]:
            if not df.empty:
                unique_offices = df["Oficina"].unique()
                assert len(unique_offices) == 1 and unique_offices[0] == DEFAULT_OFFICE_CODE, \
                    f"RN-07 violada en {name}: Oficina debe ser {DEFAULT_OFFICE_CODE}, obtenido {unique_offices}"
                unique_years = df["Año"].unique()
                assert len(unique_years) == 1 and unique_years[0] == 2026, \
                    f"RN-08 violada en {name}: Año debe ser 2026, obtenido {unique_years}"
                # Verificar formato de fechas en texto 'DD/MM/YYYY'
                for d_val in df["Desde"].tolist() + df["Hasta"].tolist():
                    assert isinstance(d_val, str) and len(d_val) == 10 and d_val[2] == "/" and d_val[5] == "/", \
                        f"Formato de fecha inválido: {d_val}"
        print(f"  [OK] RN-07: Columna 'Oficina' autocompletada con valor oficial {DEFAULT_OFFICE_CODE}")
        print(f"  [OK] RN-08: Columna 'Año' autocompletada con 2026")
        print(f"  [OK] Formato de fechas consolidado 'DD/MM/YYYY' verificado en todos los registros\n")

        # ----------------------------------------------------------------------
        # TEST 5: EXPORTACIÓN OFICIAL OPENPYXL Y FÓRMULAS
        # ----------------------------------------------------------------------
        print("--- [TEST 5] Exportación y Estilos en openpyxl ---")
        out_excel = ReportService.generate_monthly_statistics_excel(month=10, year=2026)
        assert Path(out_excel).exists(), f"El archivo Excel no fue generado: {out_excel}"

        # Cargar con openpyxl e inspeccionar hojas, celdas y fórmulas
        wb = openpyxl.load_workbook(out_excel, data_only=False)
        sheet_names = wb.sheetnames
        assert "Estadística 10-2026" in sheet_names, f"Falta hoja Estadística en {sheet_names}"
        assert "Cuadro DL-01" in sheet_names, f"Falta hoja Cuadro DL-01 en {sheet_names}"
        assert "Cuadro DL-02" in sheet_names, f"Falta hoja Cuadro DL-02 en {sheet_names}"
        print(f"  [OK] Estructura de 3 hojas oficiales generada con éxito: {sheet_names}")

        # Inspeccionar Cuadro DL-01
        ws1 = wb["Cuadro DL-01"]
        assert ws1["A1"].value.startswith("SIGET - CUADRO DL-01"), "Encabezado incorrecto en Cuadro DL-01"
        assert ws1["A4"].value == "Oficina"
        assert ws1["K4"].value == "Total Días"
        # Verificar que el DNI mantenga formato de texto
        dni_cells = [ws1.cell(row=r, column=4) for r in range(5, ws1.max_row) if ws1.cell(row=r, column=4).value]
        if dni_cells:
            assert dni_cells[0].number_format == "@", f"El DNI debe tener formato texto '@', tiene {dni_cells[0].number_format}"
        print(f"  [OK] Cuadro DL-01: Cabeceras y tipado de texto (@) en DNI validados")

        # Inspeccionar Cuadro DL-02 y fórmulas de sumatoria
        ws2 = wb["Cuadro DL-02"]
        assert ws2["A1"].value.startswith("SIGET - CUADRO DL-02"), "Encabezado incorrecto en Cuadro DL-02"
        # Última fila debe tener la fórmula SUM
        tot_cell = ws2.cell(row=ws2.max_row, column=11)
        assert str(tot_cell.value).startswith("=SUM"), f"Fórmula de total ausente en Cuadro DL-02: {tot_cell.value}"
        print(f"  [OK] Cuadro DL-02: Fórmula de administración validada: '{tot_cell.value}'")

        wb.close()
        print(f"  [OK] Archivo Excel final verificado sin corrupción: {out_excel}\n")

    finally:
        session.close()

    print("======================================================================")
    print(">> RESULTADO: [100% EXITOSO] TODAS LAS REGLAS RN-01 A RN-08 PASARON")
    print("======================================================================")


if __name__ == "__main__":
    run_all_normative_tests()
